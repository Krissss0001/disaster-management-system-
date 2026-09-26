"""Script to generate a comprehensive PowerPoint presentation (.pptx) for the Disaster Response Coordination Platform."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_NAVY = RGBColor(13, 27, 42)
    EMERGENCY_RED = RGBColor(230, 57, 70)
    TEAL_ACCENT = RGBColor(42, 157, 143)
    TEXT_MUTED = RGBColor(100, 116, 139)
    TEXT_DARK = RGBColor(15, 23, 42)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240)

    slides_data = [
        # Slide 1: Title
        {
            "type": "title",
            "title": "Disaster Response Coordination Platform",
            "subtitle": "Emergency Logistics Command: Geospatial Resource Dispatch, Demand Matching & Allocation Integrity",
            "footer": "Tech Stack: Streamlit • SQLAlchemy • SQLite • Folium • OpenStreetMap • Pytest"
        },
        # Slide 2: Problem & Motivation
        {
            "type": "standard",
            "badge": "PROBLEM & CONTEXT",
            "title": "The Emergency Coordination Challenge",
            "bullets": [
                ("The Crisis Chaos", "During natural disasters (wildfires, earthquakes, floods), relief agencies struggle with fragmented information, unverified requests, and supply mismatches."),
                ("Critical Shortages vs Unused Surpluses", "Without centralized coordination, some shelters face life-threatening deficits while neighboring staging depots hold idle inventory."),
                ("The Requirement", "A lightweight, robust, zero-configuration coordination platform that tracks incident perimeters, logs staging supplies, triages urgent field needs, and enforces mathematically sound allocation validation.")
            ]
        },
        # Slide 3: Tech Stack & Why
        {
            "type": "cards",
            "badge": "TECHNOLOGY SELECTION",
            "title": "Technology Stack: What We Used and Why",
            "cards": [
                ("Streamlit (Frontend)", "• Rapid operational interface development\n• Pure Python-native state reactivity\n• Built-in multi-page routing\n• Native integration with Pandas & Folium\n• Zero complex JavaScript build tooling"),
                ("SQLite (Database)", "• File-based, zero-server configuration\n• Highly portable (single .db file)\n• Resilient for offline or edge deployments\n• ACID-compliant transactional guarantees\n• Zero cloud or DB hosting overhead"),
                ("SQLAlchemy (ORM)", "• Declarative domain modeling with UUIDs\n• Strict foreign key & cascade enforcement\n• Decoupled from SQL dialect (Postgres-ready)\n• Scoped get_db() session transactions\n• Clean separation between DB & business logic"),
                ("Folium / OSM (Maps)", "• 100% Free OpenStreetMap & Satellite tiles\n• No API keys or credit card requirements\n• Interactive threat radius perimeters\n• Custom icon markers for needs & supplies\n• Fast client-side Leaflet rendering")
            ]
        },
        # Slide 4: System Architecture
        {
            "type": "standard",
            "badge": "SYSTEM ARCHITECTURE",
            "title": "Three-Tier Decoupled Architecture",
            "bullets": [
                ("1. Presentation Layer (Streamlit Pages)", "Consists of 6 dedicated consoles: 1_Dashboard, 2_Manage_Disasters, 3_Resources, 4_Needs, 5_Allocations, and 6_Status. Contains zero raw database queries."),
                ("2. Business Logic Layer (logic/)", "Independent pure Python modules (disasters.py, resources.py, needs.py, allocations.py, alerts.py). Houses the validate_allocation() pure validation engine and automated calling dispatch."),
                ("3. Persistence Layer (database/)", "SQLAlchemy models with UUID primary keys and scoped session management (connection.py, models.py, init_db.py). Enforces database constraints and cascade deletions.")
            ]
        },
        # Slide 5: Database Schema
        {
            "type": "cards",
            "badge": "DATA MODELING",
            "title": "SQLAlchemy Models (UUID Primary Keys)",
            "cards": [
                ("Disaster Model", "• id: UUID (Primary Key)\n• name: String (Incident title)\n• type: Wildfire, Flood, Earthquake, etc.\n• severity: Integer (Scale 1 to 5)\n• status: ongoing / contained / resolved\n• latitude, longitude, description, created_at"),
                ("Resource Model", "• id: UUID (Primary Key)\n• disaster_id: Foreign Key (Disaster)\n• type: Medical, Food/Water, Rescue, etc.\n• quantity: Total staged units\n• available: Currently unallocated units\n• provider: Red Cross, FEMA, Fire Dept\n• latitude, longitude, created_at"),
                ("ResourceNeed Model", "• id: UUID (Primary Key)\n• disaster_id: Foreign Key (Disaster)\n• type: Matches resource category\n• quantity_needed: Requested units\n• urgency: Priority scale (1 to 5)\n• status: pending / partial / fulfilled\n• latitude, longitude, created_at"),
                ("Allocation Model", "• id: UUID (Primary Key)\n• need_id: Foreign Key (ResourceNeed)\n• resource_id: Foreign Key (Resource)\n• quantity_allocated: Dispatched units\n• created_at: Immutable audit timestamp\n• Bidirectional relationship links")
            ]
        },
        # Slide 6: Core Business Logic & Pure Function
        {
            "type": "standard",
            "badge": "BUSINESS INTEGRITY",
            "title": "Allocation Validation Engine (Pure Function)",
            "bullets": [
                ("Rule 1: Strict Category Matching", "Resources can only be allocated to a need of the exact same category (e.g., Food & Water cannot fulfill Medical Supplies)."),
                ("Rule 2: Positive Quantity Constraint", "Allocation attempts with zero or negative quantities are immediately rejected with descriptive error messages."),
                ("Rule 3: Inventory Availability Cap", "Allocations cannot exceed the provider's current unallocated inventory (quantity_to_allocate <= resource.available)."),
                ("Rule 4: Demand Boundary Cap", "Prevents over-allocation past the remaining unsatisfied requirement of the need."),
                ("Rule 5: Atomic State Progression", "Decrements resource.available, logs an immutable Allocation record, and advances need.status to 'partial' or 'fulfilled'. Pure function validate_allocation() has 100% test isolation.")
            ]
        },
        # Slide 7: Emergency Calling & Alert Dispatch
        {
            "type": "standard",
            "badge": "REAL-TIME ALERTS",
            "title": "Emergency Dispatch & Calling Alert System",
            "bullets": [
                ("Automated Triggers", "Automatically triggered when high-severity disasters (Severity 4-5) or critical medical emergencies (Urgency 4-5) are posted."),
                ("Dual Mode Operation", "• Mock Mode (Default): Simulates voice call & SMS dispatch to emergency coordinators with timestamped audit logs without third-party fees.\n• Live Mode: Directly dispatches SMS/calls via Twilio REST API and posts JSON payloads to external Webhooks (PagerDuty/Slack)."),
                ("Field Visibility", "Dispatch log console in the UI lets commanders monitor outgoing calls, target telephone numbers, and delivery statuses.")
            ]
        },
        # Slide 8: Multi-Page UI Walkthrough
        {
            "type": "cards",
            "badge": "USER EXPERIENCE",
            "title": "Consolidated Incident Response Consoles",
            "cards": [
                ("1. Dashboard & Map", "• 4 real-time KPI metric cards\n• Live OpenStreetMap geospatial map\n• Incident danger perimeters\n• Searchable disaster directory"),
                ("2. Manage Disasters", "• Incident declaration form\n• Quick coordinate city presets\n• Severity slider (1 to 5)\n• Lifecycle containment status updates"),
                ("3. Resources & Needs", "• Staging intake by provider\n• Real-time availability tracking\n• Urgent shortage triage\n• Color-coded urgency badges"),
                ("4. Allocations & Status", "• Pure validation pre-check\n• Real-time mismatch prevention\n• Per-disaster surplus/deficit balance\n• Complete audit transaction ledger")
            ]
        },
        # Slide 9: Automated Testing & Verification
        {
            "type": "standard",
            "badge": "TESTING & QUALITY",
            "title": "Comprehensive Test Suite (23/23 Passing)",
            "bullets": [
                ("Phase 1 Tests", "Database schema generation, SQLite UUID primary keys, foreign key constraints, cascade deletions, and database seeding."),
                ("Phase 2-5 Tests", "Disaster queries, resource metrics, urgent need sorting, coordinate boundaries, and lifecycle status transitions."),
                ("Phase 6-7 Tests", "Comprehensive test cases for validate_allocation() covering all failure modes (type mismatch, zero/negative, over-allocation) and live balance calculations."),
                ("Execution Speed", "All 23 unit/integration tests run in ~2.3 seconds with Pytest, providing 100% confidence before deployment.")
            ]
        },
        # Slide 10: Limitations & Future Roadmap
        {
            "type": "cards",
            "badge": "ROADMAP",
            "title": "System Limitations & Concrete Next Steps",
            "cards": [
                ("Known Limitations", "• SQLite write serialization under heavy concurrency\n• Distance-agnostic supply matching (visualized but not automated)\n• Unified permissions without role-based access control (RBAC)"),
                ("Future Improvement 1: OSRM Routing", "Integrate Open Source Routing Machine (OSRM) to calculate road travel times between staging depots and incidents to auto-rank closest supplies."),
                ("Future Improvement 2: PostGIS Database", "Migrate from SQLite to PostgreSQL + PostGIS for spatial queries (ST_DWithin) and multi-agency high-concurrency writes."),
                ("Future Improvement 3: Offline PWA Sync", "Enable progressive web app offline caching for first responders in field zones with intermittent cell connectivity.")
            ]
        },
        # Slide 11: Conclusion
        {
            "type": "title",
            "title": "Thank You! Questions & Discussion",
            "subtitle": "Disaster Response Coordination Platform • Portfolio Project",
            "footer": "GitHub Repository: https://github.com/Krissss0001/disaster-management-system-"
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Background Fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        if data["type"] == "title":
            bg.fill.fore_color.rgb = DARK_NAVY
        else:
            bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

        if data["type"] == "title":
            # Title Slide
            # Top accent bar
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.8), Inches(1.2), Inches(0.08))
            bar.fill.solid()
            bar.fill.fore_color.rgb = EMERGENCY_RED
            bar.line.fill.background()

            txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(2.2))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = WHITE

            p2 = tf.add_paragraph()
            p2.text = data["subtitle"]
            p2.font.size = Pt(20)
            p2.font.color.rgb = RGBColor(203, 213, 225)
            p2.space_before = Pt(16)

            txBox_f = slide.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(10.3), Inches(0.8))
            tf_f = txBox_f.text_frame
            p_f = tf_f.paragraphs[0]
            p_f.text = data["footer"]
            p_f.font.size = Pt(14)
            p_f.font.color.rgb = TEAL_ACCENT

        elif data["type"] == "standard":
            # Header
            header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.3))
            htf = header_box.text_frame
            htf.word_wrap = True
            p_badge = htf.paragraphs[0]
            p_badge.text = data["badge"]
            p_badge.font.size = Pt(11)
            p_badge.font.bold = True
            p_badge.font.color.rgb = EMERGENCY_RED

            p_title = htf.add_paragraph()
            p_title.text = data["title"]
            p_title.font.size = Pt(26)
            p_title.font.bold = True
            p_title.font.color.rgb = TEXT_DARK
            p_title.space_before = Pt(4)

            # Bullet List
            start_y = 1.9
            for idx, (b_title, b_desc) in enumerate(data["bullets"]):
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(start_y + idx * 1.6), Inches(11.33), Inches(1.4))
                card.fill.solid()
                card.fill.fore_color.rgb = CARD_BG
                card.line.color.rgb = BORDER_COLOR
                card.line.width = Pt(1)

                # Accent line on left of card
                acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(start_y + idx * 1.6), Inches(0.1), Inches(1.4))
                acc.fill.solid()
                acc.fill.fore_color.rgb = TEAL_ACCENT if idx % 2 == 0 else EMERGENCY_RED
                acc.line.fill.background()

                c_box = slide.shapes.add_textbox(Inches(1.3), Inches(start_y + idx * 1.6 + 0.15), Inches(10.8), Inches(1.1))
                c_tf = c_box.text_frame
                c_tf.word_wrap = True
                p_bt = c_tf.paragraphs[0]
                p_bt.text = b_title
                p_bt.font.size = Pt(17)
                p_bt.font.bold = True
                p_bt.font.color.rgb = TEXT_DARK

                p_bd = c_tf.add_paragraph()
                p_bd.text = b_desc
                p_bd.font.size = Pt(13)
                p_bd.font.color.rgb = TEXT_MUTED
                p_bd.space_before = Pt(4)

        elif data["type"] == "cards":
            # Header
            header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.3), Inches(1.3))
            htf = header_box.text_frame
            htf.word_wrap = True
            p_badge = htf.paragraphs[0]
            p_badge.text = data["badge"]
            p_badge.font.size = Pt(11)
            p_badge.font.bold = True
            p_badge.font.color.rgb = EMERGENCY_RED

            p_title = htf.add_paragraph()
            p_title.text = data["title"]
            p_title.font.size = Pt(26)
            p_title.font.bold = True
            p_title.font.color.rgb = TEXT_DARK
            p_title.space_before = Pt(4)

            # 4 column cards
            card_w = 2.65
            card_gap = 0.24
            start_x = 1.0
            start_y = 2.0
            card_h = 4.8

            for idx, (c_title, c_content) in enumerate(data["cards"]):
                cx = start_x + idx * (card_w + card_gap)
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(start_y), Inches(card_w), Inches(card_h))
                card.fill.solid()
                card.fill.fore_color.rgb = CARD_BG
                card.line.color.rgb = BORDER_COLOR
                card.line.width = Pt(1)

                top_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(start_y), Inches(card_w), Inches(0.12))
                top_bar.fill.solid()
                top_bar.fill.fore_color.rgb = EMERGENCY_RED if idx == 0 else TEAL_ACCENT if idx == 1 else DARK_NAVY
                top_bar.line.fill.background()

                tx = slide.shapes.add_textbox(Inches(cx + 0.15), Inches(start_y + 0.25), Inches(card_w - 0.3), Inches(card_h - 0.4))
                tf = tx.text_frame
                tf.word_wrap = True

                pt = tf.paragraphs[0]
                pt.text = c_title
                pt.font.size = Pt(15)
                pt.font.bold = True
                pt.font.color.rgb = TEXT_DARK

                for line in c_content.split("\n"):
                    if not line.strip():
                        continue
                    pl = tf.add_paragraph()
                    pl.text = line
                    pl.font.size = Pt(11.5)
                    pl.font.color.rgb = TEXT_MUTED
                    pl.space_before = Pt(5)

    output_path = "c:/disaster/Disaster_Response_Platform_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
